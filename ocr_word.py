#!/usr/bin/env python3
"""
ocr_word.py - Image(s) -> document Word (.docx)

Chaîne de moteurs (premier qui réussit gagne) :
  1. IA de vision gratuite/en ligne : Gemini (palier gratuit) / Grok (xAI) / OpenRouter (modèles :free)
     -> les modèles sont détectés automatiquement, du plus puissant au moins puissant
  2. Repli 100 % hors ligne, sans IA générative : Tesseract (fra) puis RapidOCR
Le repli s'active automatiquement : quota dépassé, clé invalide, modèle indisponible,
erreur réseau, pas de connexion, réponse vide ou suspecte.

Dépendances : pillow, python-docx   (+ pytesseract et/ou rapidocr-onnxruntime pour le repli local)
"""

import base64
import io
import json
import os
import re
import shutil
import socket
import sys
import tempfile
import time
import unicodedata
import urllib.error
import urllib.request
from collections import Counter
from pathlib import Path

from PIL import Image, ImageOps

CONFIG_PATH = Path.home() / ".planche_ocr_config.json"
DEFAULT_CONFIG = {
    "gemini_key": "",
    "xai_key": "",
    "openrouter_key": "",
    "ordre": "gemini,xai,openrouter",
    "gemini_models": "",
    "xai_models": "",
    "openrouter_models": "",
    "tesseract_path": "",
    "langues_tesseract": "fra+eng",
    "verif_croisee": True,
    "mode_local_seul": False,
}
ENV_KEYS = {
    "gemini_key": ("GEMINI_API_KEY", "GOOGLE_API_KEY"),
    "xai_key": ("XAI_API_KEY",),
    "openrouter_key": ("OPENROUTER_API_KEY",),
}
MAX_MODELS_PER_PROVIDER = 5
QUOTA_BLOCK_SECONDS = 90

PROMPT = """You are a professional OCR engine. Transcribe ALL the text visible in this image into clean Markdown.
Rules:
- Copy the text EXACTLY as written, in its original language. Do NOT translate, correct, summarize, explain or add anything.
- Keep accents, punctuation, numbers, dates, currency symbols and special characters exactly.
- Keep the natural reading order (top to bottom; for columns, one column after the other).
- Join lines that are only visually wrapped into one paragraph. Keep distinct items on separate lines (addresses, list entries, signatures, form fields). Put a blank line between paragraphs.
- Titles/headings: use '#', '##', '###' according to their visual importance. Use **bold** / *italic* only when clearly visible.
- Bulleted lists: '- '. Numbered lists: keep the original numbers ('1. ', '2. ').
- Tables: use Markdown pipe tables with a header separator row, one row per table row, empty cells left empty.
- Handwriting: transcribe as well as you can. If a word is unreadable write [illisible]; if you are unsure, write your best guess followed by [?].
- Ignore pure decoration (photos, logos without text).
- Output ONLY the Markdown. No code fences, no introduction, no comments.
- If the image contains no text at all, output exactly: [NO_TEXT]"""


# ============================================================================
# Configuration
# ============================================================================
def load_config():
    cfg = dict(DEFAULT_CONFIG)
    try:
        cfg.update(json.loads(CONFIG_PATH.read_text(encoding="utf-8")))
    except (OSError, ValueError):
        pass
    for k, envs in ENV_KEYS.items():  # les variables d'environnement complètent
        if not str(cfg.get(k, "")).strip():
            for e in envs:
                if os.environ.get(e):
                    cfg[k] = os.environ[e]
                    break
    return cfg


def save_config(cfg):
    try:
        CONFIG_PATH.write_text(
            json.dumps(cfg, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        try:
            os.chmod(CONFIG_PATH, 0o600)
        except OSError:
            pass
        return True
    except OSError:
        return False


# ============================================================================
# Erreurs et HTTP (stdlib uniquement : pas de dépendance 'requests')
# ============================================================================
class OCRError(Exception):
    pass


class AuthError(OCRError):
    pass  # clé invalide -> on abandonne ce fournisseur


class QuotaError(OCRError):
    pass  # quota/limite -> modèle suivant


class ModelError(OCRError):
    pass  # modèle absent / réponse inutilisable -> modèle suivant


class TransientError(OCRError):
    pass  # 5xx, délai -> 1 nouvel essai puis modèle suivant


class OfflineError(OCRError):
    pass  # pas de réseau -> repli local immédiat


def _classify(code, raw):
    low = raw.lower()
    if (
        code == 429
        or "resource_exhausted" in low
        or "quota" in low
        or "rate limit" in low
        or "rate-limit" in low
    ):
        return QuotaError(f"quota/limite atteint (HTTP {code})")
    if code == 402:
        return QuotaError("crédits épuisés (HTTP 402)")
    if (
        code == 401
        or "api key not valid" in low
        or "api_key_invalid" in low
        or "incorrect api key" in low
    ):
        return AuthError(f"clé API refusée (HTTP {code})")
    if code == 403:
        return AuthError(
            "accès refusé (HTTP 403) : clé sans droit, ou région non autorisée"
        )
    if code in (400, 404, 422):
        return ModelError(f"modèle/requête refusé (HTTP {code}) : {raw[:200]}")
    if code in (408, 409, 500, 502, 503, 504, 529):
        return TransientError(f"serveur indisponible (HTTP {code})")
    return OCRError(f"erreur HTTP {code} : {raw[:200]}")


def _http(url, data=None, headers=None, timeout=120):
    body = json.dumps(data).encode("utf-8") if data is not None else None
    req = urllib.request.Request(
        url,
        data=body,
        headers=dict(headers or {}),
        method="POST" if body is not None else "GET",
    )
    if body is not None:
        req.add_header("Content-Type", "application/json")
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return json.loads(r.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        raise _classify(e.code, e.read().decode("utf-8", "replace")[:2000])
    except (socket.timeout, TimeoutError):
        raise TransientError("délai dépassé")
    except urllib.error.URLError as e:
        if isinstance(getattr(e, "reason", None), (socket.timeout, TimeoutError)):
            raise TransientError("délai dépassé")
        raise OfflineError(f"réseau indisponible ({getattr(e, 'reason', e)})")
    except (ConnectionError, OSError) as e:
        raise OfflineError(f"réseau indisponible ({e})")
    except ValueError:
        raise TransientError("réponse non lisible")


# ============================================================================
# Images
# ============================================================================
def load_image(path):
    img = ImageOps.exif_transpose(Image.open(path))
    if img.mode in ("RGBA", "LA") or (img.mode == "P" and "transparency" in img.info):
        img = img.convert("RGBA")
        bg = Image.new("RGB", img.size, "white")
        bg.paste(img, mask=img.split()[-1])
        return bg
    return img.convert("RGB")


def to_b64_jpeg(img, max_side=3000):
    img = img.convert("RGB")
    if max(img.size) > max_side:
        img = img.copy()
        img.thumbnail((max_side, max_side), Image.LANCZOS)
    for q in (92, 80, 65):  # reste sous ~15 Mo même pour de grosses images
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=q)
        if buf.tell() < 15_000_000:
            break
    return base64.b64encode(buf.getvalue()).decode("ascii")


# ============================================================================
# Fournisseurs IA (détection automatique des modèles de vision)
# ============================================================================
class Provider:
    name = "?"
    label = "?"

    def __init__(self, cfg):
        self.cfg = cfg
        self._models = None

    @property
    def key(self):
        return str(self.cfg.get(self.name + "_key", "")).strip()

    def available(self):
        return bool(self.key)

    def forced_models(self):
        return [
            m.strip()
            for m in str(self.cfg.get(self.name + "_models", "")).split(",")
            if m.strip()
        ]

    def discover(self):  # peut lever AuthError / OfflineError
        raise NotImplementedError

    def defaults(self):
        return []

    def models(self):
        if self._models is None:
            forced = self.forced_models()
            if forced:
                self._models = forced
            else:
                try:
                    found = self.discover()
                except (AuthError, OfflineError):
                    raise
                except OCRError:
                    found = []
                self._models = (found or self.defaults())[:MAX_MODELS_PER_PROVIDER]
        return self._models

    def call(self, model, b64):  # -> (texte, [avertissements])
        raise NotImplementedError


class GeminiProvider(Provider):
    name, label = "gemini", "Google Gemini"
    BASE = "https://generativelanguage.googleapis.com/v1beta"
    EXCLUDE = (
        "image",
        "tts",
        "embed",
        "live",
        "audio",
        "computer",
        "robotics",
        "imagen",
        "veo",
        "aqa",
        "deep-research",
        "customtools",
    )

    def _hdr(self):
        return {"x-goog-api-key": self.key}

    @staticmethod
    def rank(n):
        m = re.search(r"gemini-(\d+(?:\.\d+)?)", n)
        ver = float(m.group(1)) if m else 0.0
        tier = 3 if "pro" in n else (1 if "lite" in n else (2 if "flash" in n else 0))
        pre = 1 if re.search(r"preview|exp|latest", n) else 0  # stable avant préversion
        return (-ver, -tier, pre, n)

    def discover(self):
        data = _http(
            f"{self.BASE}/models?pageSize=200", headers=self._hdr(), timeout=30
        )
        names = []
        for m in data.get("models", []):
            n = m.get("name", "").replace("models/", "")
            if (
                n.startswith("gemini")
                and "generateContent" in m.get("supportedGenerationMethods", [])
                and not any(x in n for x in self.EXCLUDE)
            ):
                names.append(n)
        return sorted(names, key=self.rank)

    def defaults(self):
        return [
            "gemini-3.5-flash",
            "gemini-2.5-pro",
            "gemini-2.5-flash",
            "gemini-3.1-flash-lite",
        ]

    def call(self, model, b64):
        payload = {
            "contents": [
                {
                    "parts": [
                        {"text": PROMPT},
                        {"inline_data": {"mime_type": "image/jpeg", "data": b64}},
                    ]
                }
            ],
            "generationConfig": {"temperature": 0, "maxOutputTokens": 16384},
        }
        data = _http(
            f"{self.BASE}/models/{model}:generateContent", payload, self._hdr()
        )
        cands = data.get("candidates") or []
        if not cands:
            raise ModelError(
                "réponse vide (blocage : %s)"
                % data.get("promptFeedback", {}).get("blockReason", "inconnu")
            )
        c = cands[0]
        text = "".join(
            p.get("text", "")
            for p in c.get("content", {}).get("parts", [])
            if not p.get("thought")
        )
        fr = c.get("finishReason", "")
        if not text.strip() and fr not in ("STOP", "MAX_TOKENS", ""):
            raise ModelError(f"réponse bloquée ({fr})")
        warns = (
            ["texte possiblement tronqué (limite de longueur atteinte)"]
            if fr == "MAX_TOKENS"
            else []
        )
        return text, warns


class OpenAICompat(Provider):
    URL = ""

    def _hdr(self):
        return {"Authorization": f"Bearer {self.key}"}

    def call(self, model, b64):
        payload = {
            "model": model,
            "temperature": 0,
            "max_tokens": 8192,
            "messages": [
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": PROMPT},
                        {
                            "type": "image_url",
                            "image_url": {"url": "data:image/jpeg;base64," + b64},
                        },
                    ],
                }
            ],
        }
        data = _http(self.URL, payload, self._hdr())
        if "error" in data and not data.get("choices"):
            err = data["error"]
            code = err.get("code") if isinstance(err, dict) else None
            msg = json.dumps(err)[:300]
            try:
                raise _classify(int(code), msg)
            except (TypeError, ValueError):
                raise ModelError(f"erreur fournisseur : {msg}")
        ch = (data.get("choices") or [{}])[0]
        content = (ch.get("message") or {}).get("content") or ""
        if isinstance(content, list):
            content = "".join(p.get("text", "") for p in content if isinstance(p, dict))
        warns = (
            ["texte possiblement tronqué (limite de longueur atteinte)"]
            if ch.get("finish_reason") == "length"
            else []
        )
        return content, warns


class XAIProvider(OpenAICompat):
    name, label = "xai", "Grok (xAI)"
    URL = "https://api.x.ai/v1/chat/completions"

    def discover(self):
        data = _http(
            "https://api.x.ai/v1/language-models", headers=self._hdr(), timeout=30
        )
        out = []
        for m in data.get("models", []):
            mods = m.get("input_modalities") or []
            if "image" in mods and str(m.get("id", "")).startswith("grok"):
                out.append(m["id"])

        def rank(n):
            v = re.search(r"grok-(\d+(?:\.\d+)?)", n)
            return (
                -(float(v.group(1)) if v else 0),
                1 if re.search(r"mini|fast|lite", n) else 0,
                n,
            )

        return sorted(out, key=rank)

    def defaults(self):
        return ["grok-4", "grok-4-fast", "grok-2-vision"]


class OpenRouterProvider(OpenAICompat):
    name, label = "openrouter", "OpenRouter (modèles gratuits)"
    URL = "https://openrouter.ai/api/v1/chat/completions"

    def _hdr(self):
        h = super()._hdr()
        h["X-Title"] = "Planche photos OCR"
        return h

    def discover(self):
        data = _http("https://openrouter.ai/api/v1/models", timeout=30)
        rows = []
        for m in data.get("data", []):
            mid = str(m.get("id", ""))
            mods = (m.get("architecture") or {}).get("input_modalities") or []
            pr = m.get("pricing") or {}
            free = mid.endswith(":free") or (
                str(pr.get("prompt")) in ("0", "0.0")
                and str(pr.get("completion")) in ("0", "0.0")
            )
            if free and "image" in mods:
                b = re.search(r"(\d+(?:\.\d+)?)b\b", mid.lower())
                size = (
                    float(b.group(1))
                    if b
                    else (500 if re.search(r"gemini|gpt|grok|claude", mid) else 30)
                )
                rows.append((-size, -(m.get("context_length") or 0), mid))
        return [r[2] for r in sorted(rows)]

    def defaults(self):
        return []


PROVIDERS = {
    "gemini": GeminiProvider,
    "xai": XAIProvider,
    "openrouter": OpenRouterProvider,
}


# ============================================================================
# Repli local (sans IA générative) : Tesseract + RapidOCR
# ============================================================================
def find_tesseract(cfg):
    cands = [
        str(cfg.get("tesseract_path", "")).strip(),
        shutil.which("tesseract") or "",
    ]
    if sys.platform.startswith("win"):
        for base in (
            os.environ.get("ProgramFiles", r"C:\Program Files"),
            os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)"),
            str(Path.home() / "AppData/Local/Programs"),
        ):
            cands.append(str(Path(base) / "Tesseract-OCR" / "tesseract.exe"))
        # installation embarquée à côté du programme (PyInstaller / dossier portable)
        here = Path(getattr(sys, "_MEIPASS", Path(__file__).parent))
        cands.insert(0, str(here / "Tesseract-OCR" / "tesseract.exe"))
    for c in cands:
        if c and Path(c).is_file():
            return c
    return None


def tesseract_info(cfg):
    """-> (commande, [langues]) ou None"""
    try:
        import pytesseract
    except ImportError:
        return None
    cmd = find_tesseract(cfg)
    if not cmd:
        return None
    pytesseract.pytesseract.tesseract_cmd = cmd
    tess_dir = Path(cmd).parent / "tessdata"
    if tess_dir.is_dir():
        os.environ.setdefault("TESSDATA_PREFIX", str(tess_dir))
    try:
        return cmd, pytesseract.get_languages(config="")
    except Exception:
        return None


_rapid = None


def rapid_engine():
    global _rapid
    if _rapid is None:
        try:
            from rapidocr_onnxruntime import RapidOCR
        except ImportError:
            try:
                from rapidocr import RapidOCR
            except ImportError:
                _rapid = False
                return None
        try:
            _rapid = RapidOCR()
        except Exception:
            _rapid = False
    return _rapid or None


def _prep_local(img):
    g = ImageOps.autocontrast(ImageOps.grayscale(img), cutoff=1)
    m = max(g.size)
    if m < 1800:
        s = min(3.0, 1800 / m)
        g = g.resize((int(g.width * s), int(g.height * s)), Image.LANCZOS)
    return g


def _lines_to_text(lines):
    """lines: dicts text/left/right/top/bottom/(grp). Regroupe en paragraphes (1 ligne = 1 paragraphe Word)."""
    if not lines:
        return ""
    hs = sorted(max(1, l["bottom"] - l["top"]) for l in lines)
    med = hs[len(hs) // 2]
    minl = min(l["left"] for l in lines)
    maxr = max(l["right"] for l in lines)
    out, cur, prev = [], "", None
    for l in lines:
        if prev is not None:
            gap = l["top"] - prev["bottom"]
            short = prev["right"] < minl + 0.85 * (
                maxr - minl
            )  # ligne courte = fin de paragraphe
            brk = l.get("grp") != prev.get("grp") or gap > 0.8 * med or short
            if brk:
                out.append(cur)
                cur = l["text"]
            else:
                cur += " " + l["text"]
        else:
            cur = l["text"]
        prev = l
    out.append(cur)
    return "\n".join(t for t in out if t.strip())


def _line_mask(g):
    """Masque des traits longs (cadres, grilles de tableau) : suites de pixels sombres >= L."""
    import numpy as np

    a = (np.asarray(g) < 140).astype(np.int32)
    L = max(60, int(min(g.size) * 0.06))

    def runs(m):
        w = m.shape[1]
        if w <= L:
            return np.zeros_like(m, dtype=bool)
        cs = np.pad(np.cumsum(m, axis=1), ((0, 0), (1, 0)))
        full = (cs[:, L:] - cs[:, :-L]) == L
        f = np.pad(np.cumsum(full.astype(np.int32), axis=1), ((0, 0), (1, 0)))
        idx = np.arange(w)
        lo = np.clip(idx - L + 1, 0, full.shape[1])
        hi = np.clip(idx + 1, 0, full.shape[1])
        return (f[:, hi] - f[:, lo]) > 0

    return runs(a) | runs(a.T).T


def _tess_pass(g, langs):
    import pytesseract

    d = pytesseract.image_to_data(
        g,
        lang=langs,
        output_type=pytesseract.Output.DICT,
        config="--oem 1 --psm 3 -c preserve_interword_spaces=1",
    )
    rows, good = {}, 0
    for i, t in enumerate(d["text"]):
        t = t.strip()
        try:
            conf = float(d["conf"][i])
        except (TypeError, ValueError):
            conf = -1
        if not t or conf < 0:
            continue
        good += conf >= 50
        key = (d["block_num"][i], d["par_num"][i], d["line_num"][i])
        r = rows.setdefault(
            key,
            {
                "words": [],
                "left": 10**9,
                "right": 0,
                "top": 10**9,
                "bottom": 0,
                "grp": key[:2],
            },
        )
        r["words"].append(t)
        r["left"] = min(r["left"], d["left"][i])
        r["right"] = max(r["right"], d["left"][i] + d["width"][i])
        r["top"] = min(r["top"], d["top"][i])
        r["bottom"] = max(r["bottom"], d["top"][i] + d["height"][i])
    return [dict(r, text=" ".join(r["words"])) for r in rows.values()], good


def _ocr_tesseract(img, info, cfg):
    cmd, avail = info
    wanted = [x for x in str(cfg.get("langues_tesseract", "fra+eng")).split("+") if x]
    use = [x for x in wanted if x in avail] or (
        ["eng"] if "eng" in avail else avail[:1]
    )
    warns = []
    if "fra" in wanted and "fra" not in use:
        warns.append(
            "langue française absente de Tesseract : accents peu fiables (réinstallez avec 'French')"
        )
    langs = "+".join(use)
    g = _prep_local(img)
    lines, good = _tess_pass(g, langs)
    try:  # 2e lecture sans traits de tableau, gardée si elle lit plus de mots
        mask = _line_mask(g)
        if mask.mean() > 0.001:
            import numpy as np

            arr = np.array(g)
            arr[mask] = 255
            lines2, good2 = _tess_pass(Image.fromarray(arr), langs)
            if good2 > good:
                lines, good = lines2, good2
    except Exception:
        pass
    return _lines_to_text(lines), "Tesseract (" + langs + ")", warns


def _ocr_rapid(img, eng):
    arr = _np_array(_prep_local(img).convert("RGB"))
    res = eng(arr)
    items = []
    if isinstance(res, tuple):  # rapidocr_onnxruntime : (result, elapse)
        for box, txt, _s in res[0] or []:
            items.append((box, txt))
    else:  # rapidocr >= 2 : objet résultat
        boxes = getattr(res, "boxes", None)
        txts = getattr(res, "txts", None)
        if boxes is not None and txts is not None:
            items = list(
                zip(boxes.tolist() if hasattr(boxes, "tolist") else boxes, txts)
            )
    boxes = []
    for box, txt in items:
        xs = [p[0] for p in box]
        ys = [p[1] for p in box]
        if str(txt).strip():
            boxes.append(
                {
                    "text": str(txt).strip(),
                    "left": min(xs),
                    "right": max(xs),
                    "top": min(ys),
                    "bottom": max(ys),
                }
            )
    boxes.sort(key=lambda b: (b["top"] + b["bottom"]) / 2)
    rows = []  # fusion des boîtes situées sur la même ligne visuelle
    for b in boxes:
        h = b["bottom"] - b["top"]
        cy = (b["top"] + b["bottom"]) / 2
        if rows and abs(cy - (rows[-1]["top"] + rows[-1]["bottom"]) / 2) < 0.5 * min(
            h, rows[-1]["bottom"] - rows[-1]["top"]
        ):
            rows[-1]["parts"].append(b)
            rows[-1]["top"] = min(rows[-1]["top"], b["top"])
            rows[-1]["bottom"] = max(rows[-1]["bottom"], b["bottom"])
        else:
            rows.append({"parts": [b], "top": b["top"], "bottom": b["bottom"]})
    lines = []
    for r in rows:
        ps = sorted(r["parts"], key=lambda p: p["left"])
        txt = ps[0]["text"]
        for a, b in zip(ps, ps[1:]):
            txt += (
                "   "
                if b["left"] - a["right"] > 3 * max(1, r["bottom"] - r["top"])
                else " "
            ) + b["text"]
        lines.append(
            {
                "text": txt,
                "left": ps[0]["left"],
                "right": ps[-1]["right"],
                "top": r["top"],
                "bottom": r["bottom"],
            }
        )
    return _lines_to_text(lines), "RapidOCR", []


def _np_array(img):
    import numpy as np

    return np.array(img)


def local_ocr(img, cfg):
    """-> (texte, moteur, avertissements). Lève OCRError si aucun moteur local n'est utilisable."""
    t = tesseract_info(cfg)
    r = rapid_engine()
    plans = []
    if t and "fra" in t[1]:
        plans.append("tess")
    if r:
        plans.append("rapid")
    if t and "tess" not in plans:
        plans.append("tess")
    if not plans:
        raise OCRError(
            "Aucun moteur local disponible. Installez : pip install pytesseract rapidocr-onnxruntime "
            "et le programme Tesseract (avec la langue French)."
        )
    errs = []
    for p in plans:
        try:
            return _ocr_tesseract(img, t, cfg) if p == "tess" else _ocr_rapid(img, r)
        except Exception as e:
            errs.append(f"{p}: {e}")
    raise OCRError("Échec des moteurs locaux : " + " | ".join(errs))


# ============================================================================
# Validation, vérification croisée, pipeline
# ============================================================================
_REFUSAL = re.compile(
    r"\b(i can(?:'|no)t|i am unable|i'm unable|i'm sorry|sorry,|je ne peux pas|désolé|impossible de)\b",
    re.I,
)


def clean_ai_text(text):
    t = (text or "").strip()
    t = re.sub(r"^```[a-zA-Z]*\s*\n", "", t)
    t = re.sub(r"\n```\s*$", "", t)
    return t.strip()


def validate_ai_text(text):
    t = clean_ai_text(text)
    if not t:
        raise ModelError("réponse vide")
    if t == "[NO_TEXT]":
        return ""
    if len(t) < 250 and _REFUSAL.search(t):
        raise ModelError("le modèle a refusé / n'a pas transcrit : " + t[:80])
    return t


def _words(text):
    s = unicodedata.normalize("NFKD", text.lower())
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.findall(r"[a-z0-9]{2,}", s)


def crosscheck(a, b):
    ca, cb = Counter(_words(a)), Counter(_words(b))
    na, nb = sum(ca.values()), sum(cb.values())
    if max(na, nb) == 0:
        return 1.0
    return sum((ca & cb).values()) / max(na, nb)


class Pipeline:
    def __init__(self, cfg, log=print):
        self.cfg, self.log = cfg, log
        self.providers = []
        for n in [x.strip() for x in str(cfg.get("ordre", "")).split(",")]:
            if n in PROVIDERS:
                p = PROVIDERS[n](cfg)
                if p.available():
                    self.providers.append(p)
        self.offline = False
        self.dead_providers = set()
        self.blocked = {}  # (fournisseur, modèle) -> expiration

    def _cloud(self, img):
        b64 = to_b64_jpeg(img)
        for p in self.providers:
            if self.offline:
                return None
            if p.name in self.dead_providers:
                continue
            try:
                models = p.models()
            except AuthError as e:
                self.log(f"  ✗ {p.label} : {e} -> fournisseur ignoré")
                self.dead_providers.add(p.name)
                continue
            except OfflineError as e:
                self.log(f"  ✗ {e}")
                self.offline = True
                return None
            if not models:
                self.log(f"  ✗ {p.label} : aucun modèle de vision trouvé")
                continue
            for m in models:
                if self.blocked.get((p.name, m), 0) > time.time():
                    continue
                for attempt in (1, 2):
                    try:
                        self.log(f"  → {p.label} / {m}")
                        text, warns = p.call(m, b64)
                        text = validate_ai_text(text)
                        return {
                            "text": text,
                            "engine": f"{p.label} / {m}",
                            "warns": warns,
                        }
                    except TransientError as e:
                        self.log(
                            f"    {e}" + (" - nouvel essai" if attempt == 1 else "")
                        )
                        if attempt == 1:
                            time.sleep(2)
                            continue
                        break
                    except QuotaError as e:
                        self.log(f"    {e} -> modèle suivant")
                        self.blocked[(p.name, m)] = time.time() + QUOTA_BLOCK_SECONDS
                        break
                    except ModelError as e:
                        self.log(f"    {e} -> modèle suivant")
                        self.blocked[(p.name, m)] = float("inf")
                        break
                    except AuthError as e:
                        self.log(f"    {e} -> fournisseur ignoré")
                        self.dead_providers.add(p.name)
                        break
                    except OfflineError as e:
                        self.log(f"    {e}")
                        self.offline = True
                        return None
                if p.name in self.dead_providers:
                    break
        return None

    def run_image(self, path):
        img = load_image(path)
        res = None
        if not self.cfg.get("mode_local_seul"):
            if not self.providers:
                self.log("  (aucune clé IA configurée : mode local)")
            elif self.offline:
                self.log("  (hors ligne : mode local)")
            else:
                res = self._cloud(img)
        if res is not None:
            if self.cfg.get("verif_croisee") and res["text"]:
                try:
                    lt, le, _ = local_ocr(img, self.cfg)
                    sc = crosscheck(res["text"], lt)
                    self.log(f"  vérification croisée avec {le} : concordance {sc:.0%}")
                    if sc < 0.5:
                        res["warns"].append(
                            f"concordance faible avec l'OCR local ({sc:.0%}) : relisez attentivement"
                        )
                except OCRError:
                    pass
                except Exception:
                    pass
        else:
            if not self.cfg.get("mode_local_seul"):
                self.log("  → repli sur l'OCR local")
            text, eng, warns = local_ocr(img, self.cfg)
            res = {"text": text, "engine": eng + " [repli local]", "warns": warns}
        n_unsure = len(re.findall(r"\[illisible\]|\[\?\]", res["text"]))
        if n_unsure:
            res["warns"].append(
                f"{n_unsure} passage(s) marqué(s) [illisible] / [?] à vérifier"
            )
        if not res["text"]:
            res["warns"].append("aucun texte détecté")
        return res


# ============================================================================
# Markdown -> Word
# ============================================================================
_INLINE = re.compile(r"(\*\*[^*\n]+?\*\*|(?<![\w*])\*[^*\s][^*\n]*?\*(?![\w*]))")


def _unescape(s):
    return (
        re.sub(r"\\([*|#_\-.>`])", r"\1", s)
        .replace("<br>", "\n")
        .replace("<br/>", "\n")
    )


def _add_inline(par, text):
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = par.add_run(_unescape(part[2:-2]))
            r.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(_unescape(part[1:-1]))
            r.italic = True
        else:
            par.add_run(_unescape(part))


def _split_row(line):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|") and not s.endswith("\\|"):
        s = s[:-1]
    return [c.strip() for c in re.split(r"(?<!\\)\|", s)]


def markdown_to_docx(doc, md):
    lines = md.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1
            continue
        if ln.lstrip().startswith("|") and ln.count("|") >= 2:  # tableau
            block = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                block.append(lines[i])
                i += 1
            rows = [_split_row(b) for b in block]
            header = len(rows) > 1 and all(
                re.fullmatch(r":?-{2,}:?", c.replace(" ", "")) for c in rows[1] if c
            )
            if header:
                rows.pop(1)
            ncols = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncols)
            t.style = "Table Grid"
            for ri, r in enumerate(rows):
                for ci in range(ncols):
                    cell = t.cell(ri, ci)
                    par = cell.paragraphs[0]
                    _add_inline(par, r[ci] if ci < len(r) else "")
                    if header and ri == 0:
                        for run in par.runs:
                            run.bold = True
            doc.add_paragraph()
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", ln)
        if m:
            h = doc.add_heading(level=len(m.group(1)))
            _add_inline(h, m.group(2).strip())
            i += 1
            continue
        if re.fullmatch(r"\s*(-{3,}|\*{3,}|_{3,})\s*", ln):
            i += 1
            continue
        m = re.match(r"^(\s*)[-*•]\s+(.*)$", ln)
        if m:
            p = doc.add_paragraph(
                style="List Bullet 2" if len(m.group(1)) >= 2 else "List Bullet"
            )
            _add_inline(p, m.group(2))
            i += 1
            continue
        m = re.match(r"^(\s*)(\d+[.)])\s+(.*)$", ln)
        if m:  # numéros conservés tels quels
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = __import__("docx").shared.Cm(0.8)
            p.paragraph_format.first_line_indent = __import__("docx").shared.Cm(-0.6)
            _add_inline(p, f"{m.group(2)} {m.group(3)}")
            i += 1
            continue
        p = doc.add_paragraph()
        _add_inline(p, ln.strip())
        i += 1


def build_docx(sections, out_path):
    """sections: liste de textes Markdown (une section = une image, séparées par un saut de page)."""
    from docx import Document
    from docx.shared import Cm, Pt

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    st.paragraph_format.space_after = Pt(4)
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2)
    for n, md in enumerate(sections):
        if n:
            doc.add_page_break()
        if md.strip():
            markdown_to_docx(doc, md)
        else:
            p = doc.add_paragraph()
            r = p.add_run("(aucun texte détecté sur cette image)")
            r.italic = True
    try:
        doc.save(out_path)
    except OSError as e:
        bak = Path(tempfile.gettempdir()) / (Path(out_path).stem + "_secours.md")
        try:
            bak.write_text("\n\n---\n\n".join(sections), encoding="utf-8")
            extra = f"\nTexte sauvegardé en secours : {bak}"
        except OSError:
            extra = ""
        raise OCRError(
            f"Impossible d'écrire {out_path} ({e}). Fermez le fichier s'il est ouvert dans Word.{extra}"
        )


def process_images(paths, out_path, cfg, log=print, step=None, stop=None):
    """Retourne {'out': chemin, 'results': [(nom, moteur, avertissements)], 'ok': n, 'fail': n}"""
    pipe = Pipeline(cfg, log)
    sections, report = [], []
    for i, p in enumerate(paths, 1):
        if stop is not None and stop.is_set():
            raise OCRError("Annulé.")
        name = Path(p).name
        log(f"[{i}/{len(paths)}] {name}")
        if step:
            step(i - 1, len(paths))
        try:
            r = pipe.run_image(p)
            sections.append(r["text"])
            report.append((name, r["engine"], r["warns"]))
            log(f"  ✔ {r['engine']}" + ("".join(f"\n  ⚠ {w}" for w in r["warns"])))
        except Exception as e:
            sections.append(f"[ÉCHEC pour {name} : {e}]")
            report.append((name, "ÉCHEC", [str(e)]))
            log(f"  ✘ échec : {e}")
    build_docx(sections, out_path)
    if step:
        step(len(paths), len(paths))
    fail = sum(1 for r in report if r[1] == "ÉCHEC")
    return {
        "out": str(out_path),
        "results": report,
        "ok": len(paths) - fail,
        "fail": fail,
    }


def self_test(cfg):
    """Vérifie clés et moteurs sans consommer de quota (liste de modèles uniquement)."""
    out = []
    for n in [
        x.strip()
        for x in str(cfg.get("ordre", "")).split(",")
        if x.strip() in PROVIDERS
    ]:
        p = PROVIDERS[n](cfg)
        if not p.available():
            out.append(f"○ {p.label} : pas de clé (ignoré)")
            continue
        try:
            ms = p.models()
            out.append(
                f"✔ {p.label} : clé acceptée, {len(ms)} modèle(s) : "
                + ", ".join(ms[:3])
            )
        except OCRError as e:
            out.append(f"✘ {p.label} : {e}")
    t = tesseract_info(cfg)
    out.append(
        "✔ Tesseract : langues "
        + ", ".join(t[1])
        + ("" if "fra" in t[1] else "  (⚠ 'fra' manquant)")
        if t
        else "✘ Tesseract : introuvable (programme ou module pytesseract)"
    )
    out.append(
        "✔ RapidOCR : disponible" if rapid_engine() else "✘ RapidOCR : non installé"
    )
    return out


if __name__ == "__main__":
    # usage console : python ocr_word.py sortie.docx img1 [img2 ...] [--local]
    # (l'interface graphique est dans ocr_ui.py)
    if len(sys.argv) >= 3:
        c = load_config()
        if "--local" in sys.argv:
            c["mode_local_seul"] = True
        args = [a for a in sys.argv[1:] if a != "--local"]
        print(process_images(args[1:], args[0], c)["out"])
    else:
        print(
            "Usage : python ocr_word.py sortie.docx image1 [image2 ...] [--local]\n"
            "Interface graphique : python ocr_ui.py"
        )
